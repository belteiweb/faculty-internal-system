import os
import re
import json
import sqlite3
import openpyxl

def clean_text(text):
    if text is None:
        return ""
    return str(text).strip()

def normalize_student_id(val):
    s = clean_text(val)
    if s.isdigit():
        return str(int(s))
    return s

def split_topic_and_questions(raw_topic):
    if not raw_topic:
        return "", []
    lines = [l.strip() for l in raw_topic.splitlines() if l.strip()]
    if not lines:
        return "", []
    
    topic_title = lines[0]
    topic_title = re.sub(r'^(Topic\s*:?\s*|Research\s*Topic\s*:?\s*)', '', topic_title, flags=re.IGNORECASE).strip()
    
    questions = []
    current_q = []
    
    for line in lines[1:]:
        q_match = re.match(r'^(?:Q\d+[\.:]?|\d+[\.\)\-:])\s*(.*)', line, flags=re.IGNORECASE)
        if q_match:
            if current_q:
                questions.append(" ".join(current_q).strip())
            current_q = [q_match.group(1).strip()]
        else:
            if current_q:
                current_q.append(line)
            else:
                if '?' in line:
                    questions.append(line)
                else:
                    topic_title += " " + line
                    
    if current_q:
        questions.append(" ".join(current_q).strip())
        
    if not questions and len(lines) > 1:
        questions = lines[1:]
        
    return topic_title.strip(), questions

def parse_supervisor_info(raw_supervisor):
    if not raw_supervisor:
        return "", "", ""
    text = clean_text(raw_supervisor)
    phone_match = re.search(r'(?:Tel|Phone|Mobile|Contact)?[:\s]*([0-9\s\-]{8,15})', text, flags=re.IGNORECASE)
    phone = phone_match.group(1).replace(" ", "").replace("-", "") if phone_match else ""
    
    name_part = re.sub(r'(?:Tel|Phone|Mobile|Contact)?[:\s]*[0-9\s\-]{8,15}', '', text, flags=re.IGNORECASE).strip()
    
    # Extract title
    title_match = re.match(r'^(Mr\.|Mrs\.|Ms\.|Dr\.|Lok\s*Kru|Neak\s*Kru)', name_part, flags=re.IGNORECASE)
    title = title_match.group(1).strip() if title_match else ""
    
    clean_name = re.sub(r'^(?:Mr\.|Mrs\.|Ms\.|Dr\.|Lok\s*Kru|Neak\s*Kru)\s*', '', name_part, flags=re.IGNORECASE).strip()
    return clean_name, phone, title

def main():
    base_dir = r"D:\Project\Faculty-System"
    doc_dir = os.path.join(base_dir, "thesis-controller", "document")
    topic_file = os.path.join(doc_dir, "List", "List Supervisor and Supervisees with Topic.xlsx")
    result_file = os.path.join(doc_dir, "Result", "Thesis Result - Copy (2).xlsx")
    
    print(f"Loading topic file: {topic_file}")
    wb_topic = openpyxl.load_workbook(topic_file, data_only=True)
    print(f"Loading result file: {result_file}")
    wb_result = openpyxl.load_workbook(result_file, data_only=True)
    
    students = {}
    supervisors_set = {} # clean_name -> {name, phone, title, campus_code}
    
    campus_names = {
        "UC1": "Campus 1 (Tuol Sleng)",
        "UC2": "Campus 2 (Chom Chao)"
    }
    
    major_names = {
        "SE": "Software Engineering",
        "CNCS": "Computer Networking & Cyber Security",
        "CMD": "Computer Multimedia Design"
    }

    # Load known session student allocations from UMS files if any
    ums_known_classes = {}
    # Students in Session 22 (UC2 SE-Eve1)
    session_22_file = r"d:\1. Work\Thesis\Bachelor\Batch 6\UMS-System\session_22_manage.json"
    if os.path.exists(session_22_file):
        try:
            with open(session_22_file, "r", encoding="utf-8") as f:
                s22_data = json.load(f)
                for row in s22_data.get("students_sample", []):
                    if len(row) > 2 and row[2]:
                        s_id = str(int(row[2]))
                        ums_known_classes[s_id] = "SE-E1"
        except Exception as e:
            print("Notice on session_22_file:", e)
            
    # Students in students_tab.html
    tab_html = r"d:\1. Work\Thesis\Bachelor\Batch 6\UMS-System\students_tab.html"
    if os.path.exists(tab_html):
        try:
            from bs4 import BeautifulSoup
            with open(tab_html, "r", encoding="utf-8", errors="ignore") as f:
                soup = BeautifulSoup(f.read(), "html.parser")
                for tr in soup.select("table tbody tr"):
                    tds = tr.select("td")
                    if len(tds) >= 3 and tds[2].text.strip():
                        s_id = str(int(tds[2].text.strip()))
                        ums_known_classes[s_id] = "SE-E1"
        except Exception as e:
            pass
            
    # Class cohorts definitions per sheet
    cohort_classes = {
        "UC1-SE": ["SE-E1", "SE-E2", "SE-E3", "SE-Mor1"],
        "UC1-CNCS": ["CNCS-E1"],
        "UC1-CMD": ["CMD-E1", "CMD-E2", "CMD-E3", "CMD-E4", "CMD-SS1"],
        "UC2-SE": ["SE-E1", "SE-E2", "SE-E3", "SE-Mor1", "SE-Mor2", "SE-Aft1", "SE-SS1"],
        "UC2-CMD": ["CMD-E1", "CMD-E2", "CMD-SS1"],
        "UC2-CNCS": ["CNCS-E1", "CNCS-E2"]
    }

    # 1. Parse Topic Workbook
    for sheet_name in wb_topic.sheetnames:
        ws = wb_topic[sheet_name]
        campus_code = sheet_name.split("-")[0].strip() if "-" in sheet_name else "UC1"
        major_code = sheet_name.split("-")[1].strip() if "-" in sheet_name else "SE"
        
        current_group = None
        current_supervisor = None
        
        start_row = 1
        for r in range(1, 20):
            val_col2 = clean_text(ws.cell(r, 2).value).upper()
            if val_col2 in ["ID", "STUDENT ID"]:
                start_row = r + 1
                break
                
        sheet_students = []
        for r in range(start_row, ws.max_row + 1):
            sid_raw = ws.cell(r, 2).value
            if not sid_raw:
                continue
            sid_norm = normalize_student_id(sid_raw)
            if not sid_norm or not sid_norm.isdigit():
                continue
                
            name_latin = clean_text(ws.cell(r, 3).value)
            sex = clean_text(ws.cell(r, 4).value).upper()
            if sex not in ["M", "F"]:
                sex = "M" if sex.startswith("M") else ("F" if sex.startswith("F") else "M")
                
            topic_raw = clean_text(ws.cell(r, 5).value)
            
            group_val = ws.cell(r, 6).value
            if group_val is not None and str(group_val).strip():
                current_group = str(group_val).strip()
                
            sup_val = ws.cell(r, 7).value
            if sup_val is not None and str(sup_val).strip():
                current_supervisor = clean_text(sup_val)
                
            topic_title, questions = split_topic_and_questions(topic_raw)
            sup_name, sup_phone, sup_title = parse_supervisor_info(current_supervisor)
            
            if sup_name:
                supervisors_set[sup_name] = {
                    "name": sup_name,
                    "phone": sup_phone,
                    "title": sup_title,
                    "campus_code": campus_code
                }
                
            group_int = int(current_group) if current_group and current_group.isdigit() else None
            
            rec = {
                "student_id": sid_norm,
                "student_id_formatted": f"{int(sid_norm):08d}",
                "name_latin": name_latin,
                "sex": sex,
                "sheet": sheet_name,
                "campus_code": campus_code,
                "campus_name": campus_names.get(campus_code, f"Campus {campus_code}"),
                "major_code": major_code,
                "major_name": major_names.get(major_code, major_code),
                "group_no": group_int,
                "supervisor_name": sup_name,
                "supervisor_phone": sup_phone,
                "topic_raw": topic_raw,
                "topic_title": topic_title,
                "topic_questions": questions,
                "mark": None,
                "is_na": False,
                "status_remark": "",
                "result_status": "PENDING"
            }
            students[sid_norm] = rec
            sheet_students.append(sid_norm)
            
        # Class assignment for sheet students
        classes_available = cohort_classes.get(sheet_name, [f"{major_code}-E1"])
        num_classes = len(classes_available)
        students_per_class = max(1, len(sheet_students) // num_classes)
        
        for idx, sid in enumerate(sheet_students):
            if sid in ums_known_classes:
                assigned_class = ums_known_classes[sid]
            else:
                c_idx = min(idx // students_per_class, num_classes - 1)
                assigned_class = classes_available[c_idx]
                
            # Derive shift from class code
            if "Mor" in assigned_class or "-M" in assigned_class:
                shift = "Morning"
            elif "Aft" in assigned_class:
                shift = "Afternoon"
            elif "SS" in assigned_class:
                shift = "Weekend"
            else:
                shift = "Evening"
                
            students[sid]["class_code"] = assigned_class
            students[sid]["shift"] = shift
            students[sid]["batch"] = "Batch VI"
            students[sid]["academic_year"] = "2023-2024"

    # 2. Parse Results Workbook
    for sheet_name in wb_result.sheetnames:
        ws = wb_result[sheet_name]
        start_row = 1
        for r in range(1, 20):
            val_col2 = clean_text(ws.cell(r, 2).value).upper()
            if val_col2 in ["ID", "STUDENT ID"]:
                start_row = r + 1
                break
                
        for r in range(start_row, ws.max_row + 1):
            sid_raw = ws.cell(r, 2).value
            if not sid_raw:
                continue
            sid_norm = normalize_student_id(sid_raw)
            if not sid_norm or not sid_norm.isdigit():
                continue
                
            mark_raw = clean_text(ws.cell(r, 5).value)
            others_val = clean_text(ws.cell(r, 6).value)
            
            mark_num = None
            is_na = False
            status_remark = others_val
            result_status = "PENDING"
            
            if mark_raw:
                try:
                    mark_num = float(mark_raw)
                    if mark_num >= 60.0:
                        result_status = "PASS"
                    elif mark_num > 0.0:
                        result_status = "FAIL"
                    else:
                        result_status = "FAIL" if not others_val else others_val.upper()
                except ValueError:
                    is_na = True
                    mark_num = 0.0
                    status_remark = f"{mark_raw} {others_val}".strip()
                    result_status = mark_raw.upper()
                    
            if sid_norm in students:
                students[sid_norm]["mark"] = mark_num
                students[sid_norm]["is_na"] = is_na
                students[sid_norm]["status_remark"] = status_remark
                students[sid_norm]["result_status"] = result_status

    print(f"Total students parsed: {len(students)}")
    print(f"Total supervisors extracted: {len(supervisors_set)}")
    
    # 3. Write into SQLite Database
    db_file = os.path.join(base_dir, "database", "database.sqlite")
    conn = sqlite3.connect(db_file)
    cur = conn.cursor()
    
    # Insert Supervisors
    sup_id_map = {}
    for s_name, s_info in supervisors_set.items():
        cur.execute(
            """
            INSERT INTO supervisors (name, phone, title, campus_code, created_at, updated_at)
            VALUES (?, ?, ?, ?, datetime('now'), datetime('now'))
            """,
            (s_info["name"], s_info["phone"], s_info["title"], s_info["campus_code"])
        )
        sup_id_map[s_name] = cur.lastrowid
        
    # Insert Students
    for sid, s in students.items():
        sup_id = sup_id_map.get(s["supervisor_name"])
        cur.execute(
            """
            INSERT INTO thesis_students (
                student_id, student_id_formatted, name_latin, sex,
                campus_code, campus_name, major_code, major_name,
                class_code, shift, batch, academic_year, sheet,
                group_no, supervisor_id, supervisor_name, supervisor_phone,
                topic_title, topic_raw, topic_questions,
                score, result_status, is_na, status_remark,
                signed, created_at, updated_at
            ) VALUES (
                ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?,
                ?, ?, ?, ?,
                0, datetime('now'), datetime('now')
            )
            """,
            (
                s["student_id"], s["student_id_formatted"], s["name_latin"], s["sex"],
                s["campus_code"], s["campus_name"], s["major_code"], s["major_name"],
                s["class_code"], s["shift"], s["batch"], s["academic_year"], s["sheet"],
                s["group_no"], sup_id, s["supervisor_name"], s["supervisor_phone"],
                s["topic_title"], s["topic_raw"], json.dumps(s["topic_questions"], ensure_ascii=False),
                s["mark"], s["result_status"], 1 if s["is_na"] else 0, s["status_remark"]
            )
        )
        
    conn.commit()
    
    # Verify count
    cur.execute("SELECT COUNT(*) FROM thesis_students")
    total_st = cur.fetchone()[0]
    cur.execute("SELECT COUNT(*) FROM supervisors")
    total_sup = cur.fetchone()[0]
    print(f"DATABASE COMMITTED! Total Students in DB: {total_st}, Supervisors: {total_sup}")
    
    conn.close()

if __name__ == "__main__":
    main()
